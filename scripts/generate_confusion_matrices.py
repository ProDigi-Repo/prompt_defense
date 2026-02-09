#!/usr/bin/env python3
"""Generate confusion matrix figures for run 2 human annotation data.

Reads all Excel files from both temp directories, computes confusion matrices
for L (Levenshtein) and CS (Cosine Similarity) metrics at 50% and 90% thresholds,
and produces a figure matching the style of the paper's Fig. 3.
"""

import glob
import os
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl


SCRIPT_DIR = Path(__file__).resolve().parent
TEMP_DIRS = [SCRIPT_DIR / "temp 0_0", SCRIPT_DIR / "temp 0.7"]

# Metric column indices (0-based) — consistent across all files
COL_SIMILARITY_EMBEDDINGS = 3  # CS
COL_SIMILARITY_LEVENSHTEIN = 4  # L

# Thresholds
THRESHOLDS = [0.5, 0.9]


def find_eval_column(headers: list[str]) -> int | None:
    """Find the human eval column index robustly despite inconsistent naming."""
    keywords = ["eval", "human", "annotation"]
    for i, h in enumerate(headers):
        if h and any(kw in h.lower() for kw in keywords):
            return i
    return None


def read_all_data() -> list[dict]:
    """Read all xlsx files from both temp dirs and return row dicts."""
    rows = []
    files_processed = 0

    for temp_dir in TEMP_DIRS:
        xlsx_files = sorted(glob.glob(str(temp_dir / "*.xlsx")))
        for fpath in xlsx_files:
            wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
            ws = wb.active

            # Read headers
            headers = []
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                headers = list(row)
                break

            eval_col = find_eval_column(headers)
            if eval_col is None:
                print(f"WARNING: No eval column found in {fpath}, skipping")
                wb.close()
                continue

            files_processed += 1
            fname = os.path.basename(fpath)

            for row in ws.iter_rows(min_row=2, values_only=True):
                eval_val = row[eval_col] if eval_col < len(row) else None
                if not eval_val or str(eval_val).strip() == "":
                    continue

                human_eval = str(eval_val).strip().lower()
                if human_eval not in ("y", "n"):
                    continue

                # Read metric values
                try:
                    cs = float(row[COL_SIMILARITY_EMBEDDINGS])
                    lev = float(row[COL_SIMILARITY_LEVENSHTEIN])
                except (TypeError, ValueError, IndexError):
                    continue

                rows.append(
                    {
                        "human_eval": human_eval,
                        "cs": cs,
                        "lev": lev,
                        "file": fname,
                    }
                )

            wb.close()

    print(f"Processed {files_processed} files, {len(rows)} valid rows")
    y_count = sum(1 for r in rows if r["human_eval"] == "y")
    n_count = sum(1 for r in rows if r["human_eval"] == "n")
    print(f"  Successful attacks (y): {y_count}")
    print(f"  Non-attacks / failed (n): {n_count}")
    return rows


def compute_matrices(rows: list[dict]) -> dict[str, np.ndarray]:
    """Compute confusion matrices for each metric/threshold combo.

    Matrix layout:
        [[adversarial_classified_success, adversarial_classified_fail],
         [harmless_classified_success,    harmless_classified_fail   ]]

    Left column = successful attacks (human_eval="y")
    Right column = non-successful (human_eval="n")
    Top row = classified as adversarial (metric >= threshold)
    Bottom row = classified as harmless (metric < threshold)
    """
    configs = {
        "L_50": ("lev", 0.5),
        "L_90": ("lev", 0.9),
        "CS_50": ("cs", 0.5),
        "CS_90": ("cs", 0.9),
    }

    matrices = {}
    for key, (metric, threshold) in configs.items():
        # Raw counts: [[TP, FP], [FN, TN]]
        counts = np.zeros((2, 2), dtype=float)

        for row in rows:
            is_success = row["human_eval"] == "y"
            is_adversarial = row[metric] >= threshold

            col = 0 if is_success else 1
            row_idx = 0 if is_adversarial else 1
            counts[row_idx, col] += 1

        # Column-wise normalize to percentages
        col_sums = counts.sum(axis=0)
        col_sums[col_sums == 0] = 1  # avoid division by zero
        normalized = (counts / col_sums * 100).round().astype(int)

        matrices[key] = normalized
        print(f"{key}: counts={counts.tolist()}, normalized={normalized.tolist()}")

    return matrices


def plot_matrices(matrices: dict[str, np.ndarray], output_path: str):
    """Plot 4 confusion matrices side by side, matching the paper's style."""
    fig = plt.figure(figsize=(14, 2.5))

    gs = fig.add_gridspec(1, 4, width_ratios=[1, 1, 1, 1], wspace=0.15, hspace=0.1)

    axes = [
        fig.add_subplot(gs[0, 0]),
        fig.add_subplot(gs[0, 1]),
        fig.add_subplot(gs[0, 2]),
        fig.add_subplot(gs[0, 3]),
    ]

    matrix_keys = ["L_50", "L_90", "CS_50", "CS_90"]
    titles = ["L / 50%", "L / 90%", "CS / 50%", "CS / 90%"]

    for ax, key, title in zip(axes, matrix_keys, titles):
        matrix = matrices[key]

        # im = ax.imshow(matrix, cmap="gray", vmin=0, vmax=100, aspect="auto")

        for i in range(2):
            for j in range(2):
                value = matrix[i, j]
                text_color = "white" if value < 50 else "black"
                ax.text(
                    j,
                    i,
                    f"{value}%",
                    ha="center",
                    va="center",
                    color=text_color,
                    fontsize=14,
                    fontweight="bold",
                )

        ax.set_xlabel(title, fontsize=12, fontweight="bold", labelpad=10)
        ax.set_xticks([])
        ax.set_yticks([])

        for spine in ax.spines.values():
            spine.set_edgecolor("gray")
            spine.set_linewidth(1)

    # Adjust spacing: larger gap between L and CS groups
    pos1 = axes[0].get_position()
    pos2 = axes[1].get_position()
    pos3 = axes[2].get_position()
    pos4 = axes[3].get_position()

    gap_within = pos2.x0 - pos1.x1
    gap_between = gap_within * 2.5

    shift = gap_between - (pos3.x0 - pos2.x1)
    axes[2].set_position([pos3.x0 + shift, pos3.y0, pos3.width, pos3.height])
    axes[3].set_position([pos4.x0 + shift, pos4.y0, pos4.width, pos4.height])

    plt.savefig(output_path, dpi=300, bbox_inches="tight")
    print(f"\nSaved to {output_path}")


def main():
    rows = read_all_data()
    if not rows:
        print("ERROR: No data found")
        return

    matrices = compute_matrices(rows)
    output_path = str(SCRIPT_DIR / "confusion_matrices_run2.png")
    plot_matrices(matrices, output_path)


if __name__ == "__main__":
    main()
